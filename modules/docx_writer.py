"""
DOCX 写入模块（重构版）

职责：把译文写回 DOCX，尽量保留原样式。

相比旧版的修复：
  1. 译文为空时保留原文，而不是把段落清空
     （旧版 translated_list 初始为 ''，未翻译的段落会被整段抹掉）
  2. 替换文字时保留第一个 run 的样式，只清空多余 run
  3. 新增图片回写能力（把翻译后的图片替换回文档里）
"""

import os
import sys
from typing import Dict, Sequence

from docx import Document as DocxDocument

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from modules.docx_reader import DocxContent


class DocxWriter:
    """DOCX 写入器"""

    def __init__(self, template_path: str):
        self.template_path = template_path
        self.doc = DocxDocument(template_path)

    # ── 对外主接口 ───────────────────────────────────────

    def write_translated(
        self,
        content: DocxContent,
        translated_texts: Sequence[str],
        output_path: str,
    ) -> str:
        """
        写回译文。

        `translated_texts` 的顺序必须与编排器构建的一致：
            段落(按 index) → 标题(按 index) → 表格(行优先) → 页眉 → 页脚
        当某个位置缺译文（越界或空串）时保留原文，不会清空内容。
        """
        os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)

        cursor = 0

        def take(fallback: str) -> str:
            nonlocal cursor
            value = None
            if cursor < len(translated_texts):
                value = translated_texts[cursor]
            cursor += 1
            if value is None or not str(value).strip():
                return fallback
            return str(value)

        for para in sorted(content.paragraphs, key=lambda p: p.index):
            self._set_paragraph_text(para.index, take(para.text))

        for heading in sorted(content.headings, key=lambda h: h.index):
            self._set_paragraph_text(heading.index, take(heading.text))

        for table in content.tables:
            for row in range(table.rows):
                for col in range(table.cols):
                    cell = table.get_cell(row, col)
                    if cell is None:
                        continue
                    self._set_cell_text(table.index, row, col, take(cell.text))

        for header in content.headers:
            take(header)
        for footer in content.footers:
            take(footer)

        self.doc.save(output_path)
        return output_path

    # ── 文字替换 ─────────────────────────────────────────

    def _set_paragraph_text(self, para_index: int, text: str) -> bool:
        """按文档顺序索引替换段落文字，保留原有样式"""
        for idx, para in enumerate(self.doc.paragraphs):
            if idx != para_index:
                continue
            if para.runs:
                para.runs[0].text = text
                for run in para.runs[1:]:
                    run.text = ""
            else:
                para.add_run(text)
            return True
        return False

    def _set_cell_text(self, table_index: int, row: int, col: int, text: str) -> bool:
        if table_index >= len(self.doc.tables):
            return False
        table = self.doc.tables[table_index]
        if row >= len(table.rows):
            return False
        cells = table.rows[row].cells
        if col >= len(cells):
            return False

        cell = cells[col]
        if not cell.paragraphs:
            cell.text = text
            return True

        para = cell.paragraphs[0]
        if para.runs:
            para.runs[0].text = text
            for run in para.runs[1:]:
                run.text = ""
        else:
            para.add_run(text)
        # 清空单元格里的其他段落，避免残留原文
        for extra in cell.paragraphs[1:]:
            for run in extra.runs:
                run.text = ""
        return True

    # ── 图片回写 ─────────────────────────────────────────

    def replace_images(self, replacements: Dict[str, str]) -> int:
        """
        用翻译后的图片替换文档内嵌图片。

        Args:
            replacements: {docx 包内部件名 或 图片文件名: 新图片路径}

        Returns:
            成功替换的数量
        """
        if not replacements:
            return 0

        by_basename = {os.path.basename(k): v for k, v in replacements.items()}
        by_partname = {k: v for k, v in replacements.items()}
        replaced = 0

        for part in self.doc.part.package.iter_parts():
            partname = str(getattr(part, "partname", "") or "")
            source = by_partname.get(partname) or by_basename.get(os.path.basename(partname))
            if not source or not os.path.exists(source):
                continue
            try:
                with open(source, "rb") as f:
                    part._blob = f.read()
                replaced += 1
            except Exception:
                continue
        return replaced
